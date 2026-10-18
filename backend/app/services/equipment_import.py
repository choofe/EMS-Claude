"""
Excel/CSV equipment import (spec s.21), per-group, preview -> confirm.

Workflow: the manager picks ONE group, uploads a file in the standard format (see import_parser) and gets a
preview: how many rows are valid, and for every rejected row its row number, the offending value and the
reason. Confirming imports the VALID rows only, in ONE transaction with ONE audit row; the rejected rows
are reported back, never silently fixed. Nothing is ever updated or overwritten: an equipment code that
already exists (any letter case, any group) is rejected as "already_exists".

The preview is stateless: confirm re-uploads the same file with the preview's SHA-256 and the server
validates everything again — a stale or tampered preview cannot import anything the server would not accept.

Safety net for mistakes: every import is a batch (equipment.import_batch_id). undo_batch deletes the whole
batch — the single exception to "never hard-delete", allowed only while NO equipment of the batch has
history (no report of any kind, active or soft-deleted), nothing was moved to another group, and only for
Management. The batch row and an audit row (with the deleted codes) remain.
"""
import asyncio
import hashlib
import os
import re
from dataclasses import dataclass
from datetime import datetime, timezone

from sqlalchemy import delete, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import DomainError
from app.core.normalize import canonical_equipment_code, clean_text
from app.models.equipment import Equipment
from app.models.group import Group
from app.models.import_batch import EquipmentImportBatch
from app.models.report import Report
from app.models.user import User
from app.services.audit import add_audit
from app.services.import_parser import RawRow, parse_upload
from app.services.queries import lock_group

_CODE_CHARS = re.compile(r"[A-Za-z0-9._/-]")
_DB_CHUNK = 500
_SAMPLE = 10


@dataclass
class ValidRow:
    row: int
    code: str
    description: str | None


def safe_file_name(name: str | None) -> str:
    base = os.path.basename((name or "").replace("\\", "/"))
    base = "".join(ch for ch in base if ch.isprintable())[:255].strip()
    return base or "upload"


def _shown(ch: str) -> str:
    return ch if ch.isprintable() and not ch.isspace() else f"U+{ord(ch):04X}"


def _bad_chars(value: str) -> list[str]:
    bad: list[str] = []
    for ch in value:
        if not _CODE_CHARS.fullmatch(ch):
            shown = _shown(ch)
            if shown not in bad:
                bad.append(shown)
        if len(bad) >= 5:
            break
    return bad


def _error(row: int, value: str, code: str, **extra) -> dict:
    return {"row": row, "equipment_code": value.strip()[:80], "code": code, **extra}


def _validate_row(raw: RawRow, group: Group) -> ValidRow | dict:
    code_cell, group_cell = raw.cells["equipment_code"], raw.cells["group_code"]
    desc_cell = raw.cells.get("description")
    shown = code_cell.text
    for cell in (code_cell, group_cell, desc_cell):
        if cell is not None and cell.issue == "formula":
            return _error(raw.row, shown, "formula_not_allowed")
    if code_cell.issue == "date" or group_cell.issue == "date":
        return _error(raw.row, shown, "date_cell", column="equipment_code" if code_cell.issue == "date" else "group_code")
    if not code_cell.text.strip():
        return _error(raw.row, shown, "missing_equipment_code")
    if not group_cell.text.strip():
        return _error(raw.row, shown, "missing_group_code")
    try:
        code = canonical_equipment_code(code_cell.text)
    except DomainError:
        return _error(raw.row, shown, "invalid_equipment_code", bad_chars=_bad_chars(code_cell.text.strip()))
    if group_cell.text.strip().upper() != group.code:
        return _error(raw.row, shown, "group_mismatch", found=group_cell.text.strip()[:32], expected=group.code)
    description = None
    if desc_cell is not None and desc_cell.text.strip():
        try:
            description = clean_text(desc_cell.text, max_length=2000, code="invalid_description")
        except DomainError:
            return _error(raw.row, shown, "invalid_description")
    return ValidRow(raw.row, code, description)


async def _existing(db: AsyncSession, codes: list[str]) -> dict[str, tuple[str, bool]]:
    found: dict[str, tuple[str, bool]] = {}
    for i in range(0, len(codes), _DB_CHUNK):
        chunk = codes[i:i + _DB_CHUNK]
        rows = (await db.execute(
            select(Equipment.equipment_code, Group.code, Equipment.is_active)
            .join(Group, Group.id == Equipment.group_id).where(Equipment.equipment_code.in_(chunk))
        )).all()
        found.update({c: (g, a) for c, g, a in rows})
    return found


async def _active_group(db: AsyncSession, group_id: int, *, lock: bool) -> Group:
    if lock:
        group = await lock_group(db, group_id, exclusive=False)
    else:
        group = (await db.execute(select(Group).where(Group.id == group_id))).scalar_one_or_none()
        if group is None:
            raise DomainError("group_not_found", 404, group_id=group_id)
    if not group.is_active:
        raise DomainError("group_inactive", 409, group_id=group_id)
    return group


async def _analyse(db: AsyncSession, group: Group, filename: str, data: bytes) -> tuple[dict, list[ValidRow]]:
    # CPU-bound (zip + XML): keep it off the event loop so one big upload cannot stall every other request.
    parsed = await asyncio.to_thread(parse_upload, filename, data)
    valid: list[ValidRow] = []
    errors: list[dict] = []
    first_seen: dict[str, int] = {}
    for raw in parsed.rows:
        result = _validate_row(raw, group)
        if isinstance(result, dict):
            errors.append(result)
        elif result.code in first_seen:
            errors.append(_error(raw.row, raw.cells["equipment_code"].text, "duplicate_in_file", first_row=first_seen[result.code]))
        else:
            first_seen[result.code] = result.row
            valid.append(result)
    taken = await _existing(db, [v.code for v in valid])
    if taken:
        still_valid = []
        for v in valid:
            if v.code in taken:
                existing_group, is_active = taken[v.code]
                errors.append(_error(v.row, v.code, "already_exists", existing_group_code=existing_group,
                                     existing_is_active=is_active))
            else:
                still_valid.append(v)
        valid = still_valid
    errors.sort(key=lambda e: e["row"])
    report = {
        "group": {"id": group.id, "code": group.code, "name": group.name},
        "file_name": safe_file_name(filename),
        "file_sha256": hashlib.sha256(data).hexdigest(),
        "rows_total": len(parsed.rows),
        "valid_count": len(valid),
        "error_count": len(errors),
        "ignored_columns": parsed.ignored_columns,
        "errors": errors,
        "sample": [{"row": v.row, "equipment_code": v.code, "description": v.description} for v in valid[:_SAMPLE]],
    }
    return report, valid


async def preview_import(db: AsyncSession, *, group_id: int, filename: str, data: bytes) -> dict:
    group = await _active_group(db, group_id, lock=False)
    report, _ = await _analyse(db, group, filename, data)
    return report


async def commit_import(db: AsyncSession, *, actor_id: int, group_id: int, filename: str, data: bytes,
                        expected_sha256: str) -> dict:
    if hashlib.sha256(data).hexdigest() != (expected_sha256 or "").strip().lower():
        raise DomainError("file_changed", 409)  # not the file that was previewed
    group = await _active_group(db, group_id, lock=True)
    report, valid = await _analyse(db, group, filename, data)
    if not valid:
        raise DomainError("nothing_to_import", 422, error_count=report["error_count"])
    batch = EquipmentImportBatch(
        group_id=group.id, file_name=report["file_name"], file_sha256=report["file_sha256"], created_by=actor_id,
        created_count=len(valid), rejected_count=report["error_count"],
    )
    db.add(batch)
    await db.flush()
    db.add_all(Equipment(equipment_code=v.code, group_id=group.id, description=v.description,
                         import_batch_id=batch.id) for v in valid)
    try:
        await db.flush()
    except IntegrityError:  # someone created one of these codes between our check and the insert
        await db.rollback()
        raise DomainError("import_conflict", 409)
    add_audit(db, actor_id=actor_id, action="equipment.import", entity_type="equipment_import_batch",
              entity_id=batch.id, metadata={"group_id": group.id, "group_code": group.code,
                                            "file_name": report["file_name"], "file_sha256": report["file_sha256"],
                                            "rows_total": report["rows_total"], "created": len(valid),
                                            "rejected": report["error_count"]})
    await db.commit()
    return {**report, "batch_id": batch.id, "created_count": len(valid)}


# ------------------------------------------------------------------ batches / undo
async def _batch_stats(db: AsyncSession, batch_ids: list[int]) -> dict[int, dict]:
    stats = {i: {"remaining": 0, "reports": 0, "moved": 0} for i in batch_ids}
    if not batch_ids:
        return stats
    for bid, n in (await db.execute(select(Equipment.import_batch_id, func.count(Equipment.id))
                                    .where(Equipment.import_batch_id.in_(batch_ids))
                                    .group_by(Equipment.import_batch_id))).all():
        stats[bid]["remaining"] = n
    for bid, n in (await db.execute(select(Equipment.import_batch_id, func.count(Report.id))
                                    .join(Report, Report.equipment_id == Equipment.id)
                                    .where(Equipment.import_batch_id.in_(batch_ids))
                                    .group_by(Equipment.import_batch_id))).all():
        stats[bid]["reports"] = n
    for bid, n in (await db.execute(select(Equipment.import_batch_id, func.count(Equipment.id))
                                    .join(EquipmentImportBatch, EquipmentImportBatch.id == Equipment.import_batch_id)
                                    .where(Equipment.import_batch_id.in_(batch_ids),
                                           Equipment.group_id != EquipmentImportBatch.group_id)
                                    .group_by(Equipment.import_batch_id))).all():
        stats[bid]["moved"] = n
    return stats


async def list_batches(db: AsyncSession, *, limit: int, offset: int) -> tuple[list[dict], int]:
    total = (await db.execute(select(func.count(EquipmentImportBatch.id)))).scalar_one()
    rows = (await db.execute(
        select(EquipmentImportBatch, Group.code, Group.name, User.full_name)
        .join(Group, Group.id == EquipmentImportBatch.group_id)
        .join(User, User.id == EquipmentImportBatch.created_by)
        .order_by(EquipmentImportBatch.id.desc()).limit(limit).offset(offset)
    )).all()
    stats = await _batch_stats(db, [b.id for b, *_ in rows])
    items = []
    for b, gcode, gname, creator in rows:
        s = stats[b.id]
        items.append({
            "id": b.id, "group_id": b.group_id, "group_code": gcode, "group_name": gname, "file_name": b.file_name,
            "created_by_name": creator, "created_at": b.created_at, "created_count": b.created_count,
            "rejected_count": b.rejected_count, "remaining_count": s["remaining"], "report_count": s["reports"],
            "undone_at": b.undone_at,
            "can_undo": b.undone_at is None and s["remaining"] > 0 and s["reports"] == 0 and s["moved"] == 0,
        })
    return items, total


async def _batch_has_reports(db: AsyncSession, batch_id: int) -> DomainError:
    """One response contract for 'cannot undo: reports exist', whether the guard or the FK caught it."""
    report_count = (await _batch_stats(db, [batch_id]))[batch_id]["reports"]
    codes = (await db.execute(
        select(Equipment.equipment_code).join(Report, Report.equipment_id == Equipment.id)
        .where(Equipment.import_batch_id == batch_id).distinct().order_by(Equipment.equipment_code).limit(20)
    )).scalars().all()
    return DomainError("batch_has_reports", 409, report_count=report_count, equipment_codes=list(codes))


async def undo_batch(db: AsyncSession, *, actor_id: int, batch_id: int) -> dict:
    batch = (await db.execute(select(EquipmentImportBatch).where(EquipmentImportBatch.id == batch_id)
                              .with_for_update().execution_options(populate_existing=True))).scalar_one_or_none()
    if batch is None:
        raise DomainError("import_batch_not_found", 404)
    if batch.undone_at is not None:
        raise DomainError("batch_already_undone", 409)
    equipment = (await db.execute(select(Equipment).where(Equipment.import_batch_id == batch_id)
                                  .order_by(Equipment.equipment_code).with_for_update()
                                  .execution_options(populate_existing=True))).scalars().all()
    if not equipment:
        raise DomainError("batch_empty", 409)
    stats = (await _batch_stats(db, [batch_id]))[batch_id]
    if stats["reports"]:
        raise await _batch_has_reports(db, batch_id)
    if stats["moved"]:
        raise DomainError("batch_equipment_moved", 409, moved_count=stats["moved"])
    codes = [e.equipment_code for e in equipment]
    try:
        await db.execute(delete(Equipment).where(Equipment.import_batch_id == batch_id))
        batch.undone_at = datetime.now(timezone.utc)
        batch.undone_by = actor_id
        add_audit(db, actor_id=actor_id, action="equipment.import_undo", entity_type="equipment_import_batch",
                  entity_id=batch.id, metadata={"group_id": batch.group_id, "deleted_count": len(codes), "codes": codes})
        await db.commit()
    except IntegrityError:  # a report referencing one of them appeared concurrently: the FK is the last line of defence
        await db.rollback()
        raise await _batch_has_reports(db, batch_id)
    return {"batch_id": batch_id, "deleted_count": len(codes)}


# ------------------------------------------------------------------ template
def build_template(fmt: str) -> tuple[bytes, str, str]:
    """(content, media type, file name). Header only — no sample row, so nothing can be imported by accident."""
    if fmt == "csv":
        return "\ufeffequipment_code,group_code,description\r\n".encode("utf-8"), "text/csv; charset=utf-8", \
            "equipment-import-template.csv"
    import io
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    wb = Workbook()
    ws = wb.active
    ws.title = "Equipment"
    ws.append(["equipment_code", "group_code", "description"])
    notes = ("Required. Letters, digits and . _ / - only; no spaces. Must not exist yet.",
             "Required. Must equal the group you import into (e.g. ELV).", "Optional.")
    for col, note in zip("ABC", notes):
        ws[f"{col}1"].comment = Comment(note, "EMS")
        ws.column_dimensions[col].width = 28 if col != "C" else 48
        for row in range(2, 1002):
            ws[f"{col}{row}"].number_format = "@"  # text: stops Excel turning codes like 1-10 into dates
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", \
        "equipment-import-template.xlsx"
