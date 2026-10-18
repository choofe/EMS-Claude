"""
Reads an uploaded .xlsx / .csv into raw rows. No database, no business rules (those live in
equipment_import.py) — this module only answers "is this a safe, well-shaped single table, and what
are its cells?". File-level problems raise DomainError (nothing is imported); row-level problems are
left to the caller.

Defences (spec s.21/s.27): extension + magic-byte check; hard size / row limits; xlsx is opened as a zip
FIRST and rejected for macros, too many members, absurd compression ratios or a huge uncompressed
total (zip bomb) before openpyxl ever sees it; openpyxl runs read-only with defusedxml installed
(XML entity bombs / XXE are refused); formulas are never evaluated — a cell starting with "=" is handed
back as text and the row is rejected by the caller; date-typed cells are flagged because Excel silently
turns codes like "1-10" into dates.

Standard format (one sheet, header in row 1): equipment_code, group_code, description (optional).
"""
import csv
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time

from openpyxl import load_workbook

from app.core.config import get_settings
from app.core.errors import DomainError

REQUIRED_COLUMNS = ("equipment_code", "group_code")
OPTIONAL_COLUMNS = ("description",)
_KNOWN = REQUIRED_COLUMNS + OPTIONAL_COLUMNS
_MAX_PHYSICAL_ROWS = 50_000   # a sheet can declare a million blank rows; never walk more than this
_MAX_ZIP_MEMBERS = 200
_MAX_COMPRESSION_RATIO = 200
_RATIO_CHECK_MIN_BYTES = 64 * 1024
# Members that mean macros / active content. NOT every ".bin": Excel itself writes
# xl/printerSettings/printerSettings1.bin into ordinary, harmless workbooks.
_FORBIDDEN_MEMBERS = ("vbaproject", "xl/activex/", "xl/embeddings/")


@dataclass
class RawCell:
    text: str = ""
    issue: str | None = None   # "formula" | "date"


@dataclass
class RawRow:
    row: int                    # 1-based row number as the user sees it in Excel (header = 1)
    cells: dict[str, RawCell] = field(default_factory=dict)


@dataclass
class ParsedFile:
    rows: list[RawRow]
    ignored_columns: list[str]


def _cell(value) -> RawCell:
    if value is None:
        return RawCell()
    if isinstance(value, (datetime, date, time)):
        return RawCell(str(value), "date")
    if isinstance(value, bool):
        return RawCell(str(value).upper())
    if isinstance(value, float) and value.is_integer():
        return RawCell(str(int(value)))
    text = str(value)
    if text.lstrip().startswith("="):
        return RawCell(text, "formula")
    return RawCell(text)


def _header_key(value) -> str:
    return "" if value is None else str(value).replace("\xa0", " ").strip().lower()


def _map_header(header_values) -> tuple[dict[int, str], list[str]]:
    """column index -> standard name; plus the non-empty headers we do not use."""
    mapping: dict[int, str] = {}
    ignored: list[str] = []
    seen: set[str] = set()
    for idx, raw in enumerate(header_values):
        key = _header_key(raw)
        if not key:
            continue
        if key in _KNOWN:
            if key in seen:
                raise DomainError("duplicate_columns", 422, column=key)
            seen.add(key)
            mapping[idx] = key
        else:
            ignored.append(str(raw).strip()[:64])
    missing = [c for c in REQUIRED_COLUMNS if c not in seen]
    if missing:
        raise DomainError("missing_columns", 422, columns=missing)
    return mapping, ignored


def _collect(rows_iter, max_rows: int) -> ParsedFile:
    """rows_iter yields tuples of raw cell values, starting with the header row."""
    header = next(rows_iter, None)
    if header is None or all(_header_key(v) == "" for v in header):
        raise DomainError("empty_file", 422)
    mapping, ignored = _map_header(header)
    out: list[RawRow] = []
    for number, values in enumerate(rows_iter, start=2):
        if number > _MAX_PHYSICAL_ROWS:
            raise DomainError("too_many_rows", 422, max_rows=max_rows)
        cells = {name: _cell(values[idx] if idx < len(values) else None) for idx, name in mapping.items()}
        if all(c.text.strip() == "" and c.issue is None for c in cells.values()):
            continue  # fully blank row
        if len(out) >= max_rows:
            raise DomainError("too_many_rows", 422, max_rows=max_rows)
        out.append(RawRow(row=number, cells=cells))
    if not out:
        raise DomainError("no_data_rows", 422)
    return ParsedFile(rows=out, ignored_columns=ignored)


# ---------------------------------------------------------------- xlsx
def _check_zip_safety(data: bytes, max_uncompressed: int) -> None:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except (zipfile.BadZipFile, ValueError, OSError):
        raise DomainError("file_corrupt", 422)
    if len(infos) > _MAX_ZIP_MEMBERS:
        raise DomainError("file_unsafe", 422)
    names = {i.filename.lower() for i in infos}
    if "[content_types].xml" not in names or not any(n.startswith("xl/worksheets/") for n in names):
        raise DomainError("file_corrupt", 422)
    if any(marker in n for n in names for marker in _FORBIDDEN_MEMBERS):
        raise DomainError("file_unsafe", 422)  # macros / active content / embedded objects
    total = 0
    for info in infos:
        total += info.file_size
        if info.file_size > 0 and info.compress_size > 0 and info.file_size / info.compress_size > _MAX_COMPRESSION_RATIO \
                and info.file_size > _RATIO_CHECK_MIN_BYTES:
            raise DomainError("file_unsafe", 422)
        if total > max_uncompressed:
            raise DomainError("file_unsafe", 422)


def _parse_xlsx(data: bytes, max_rows: int, max_uncompressed: int) -> ParsedFile:
    if not data.startswith(b"PK\x03\x04"):
        raise DomainError("file_corrupt", 422)
    _check_zip_safety(data, max_uncompressed)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except DomainError:
        raise
    except Exception:  # openpyxl/defusedxml raise many types for malformed or hostile files
        raise DomainError("file_corrupt", 422)
    try:
        if len(wb.sheetnames) != 1:
            raise DomainError("multiple_sheets", 422, sheets=len(wb.sheetnames))
        ws = wb[wb.sheetnames[0]]
        try:
            return _collect(ws.iter_rows(values_only=True), max_rows)
        except DomainError:
            raise
        except Exception:
            raise DomainError("file_corrupt", 422)
    finally:
        wb.close()


# ---------------------------------------------------------------- csv
def _parse_csv(data: bytes, max_rows: int) -> ParsedFile:
    if b"\x00" in data:
        raise DomainError("file_corrupt", 422)
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise DomainError("file_encoding", 422)  # ask the user to save as "CSV UTF-8"
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    try:
        return _collect(csv.reader(io.StringIO(text, newline=""), dialect), max_rows)
    except csv.Error:
        raise DomainError("file_corrupt", 422)


def parse_upload(filename: str, data: bytes) -> ParsedFile:
    settings = get_settings()
    if len(data) > settings.import_max_bytes:
        raise DomainError("file_too_large", 413, max_bytes=settings.import_max_bytes)
    if not data:
        raise DomainError("empty_file", 422)
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return _parse_xlsx(data, settings.import_max_rows, settings.import_max_uncompressed_bytes)
    if name.endswith(".csv"):
        return _parse_csv(data, settings.import_max_rows)
    raise DomainError("unsupported_file_type", 422)  # .xls / .xlsm / anything else
