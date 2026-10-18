"""Phase 5 — equipment import (preview / commit / undo / template). Runs on SQLite with the REAL routers and
the real login flow; PostgreSQL-specific behaviour is in tests/postgres/test_import_pg.py."""
import hashlib
import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

import app.models as m
from app.core.config import get_settings
from tests import factories as f

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER = ["equipment_code", "group_code", "description"]
NBSP = "\xa0"


def make_xlsx(rows, header=HEADER, sheets=1):
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    if header is not None:
        ws.append(header)
    for r in rows:
        ws.append(r)
    for i in range(2, sheets + 1):
        wb.create_sheet(f"S{i}").append(header or [])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


async def preview(api, who, group_id, data, name="a.xlsx"):
    return await api.post("/equipment/import/preview", who, data={"group_id": str(group_id)},
                          files={"file": (name, data, XLSX)})


async def commit(api, who, group_id, data, name="a.xlsx", digest=None):
    return await api.post("/equipment/import/commit", who,
                          data={"group_id": str(group_id), "preview_sha256": digest or sha(data)},
                          files={"file": (name, data, XLSX)})


def reasons(resp):
    return {e["row"]: e["code"] for e in resp.json()["errors"]}


async def count(db, model):
    return (await db.execute(select(func.count()).select_from(model))).scalar_one()


# ======================= happy path & canonical forms =======================

async def test_preview_then_commit_imports_valid_rows_with_canonical_codes(api, world, db):
    data = make_xlsx([
        [f"se-01-1.1{NBSP}", "ELV", None],            # trailing NBSP + lower case (real-world file quirk)
        ["SE-01-1.2", " elv ", "  main   lift "],
        ["SE-02-1", "ELV", ""],
    ])
    p = await preview(api, "boss", world.elv.id, data)
    assert p.status_code == 200, p.text
    body = p.json()
    assert (body["rows_total"], body["valid_count"], body["error_count"]) == (3, 3, 0)
    assert body["group"] == {"id": world.elv.id, "code": "ELV", "name": "Elevator"}
    assert body["file_sha256"] == sha(data)
    assert await count(db, m.Equipment) == 2  # preview wrote nothing (ELV-1, DOOR-1 from the world fixture)

    c = await commit(api, "boss", world.elv.id, data)
    assert c.status_code == 201, c.text
    out = c.json()
    assert out["created_count"] == 3 and out["batch_id"]
    rows = (await db.execute(select(m.Equipment).where(m.Equipment.import_batch_id == out["batch_id"])
                             .order_by(m.Equipment.equipment_code))).scalars().all()
    assert [e.equipment_code for e in rows] == ["SE-01-1.1", "SE-01-1.2", "SE-02-1"]
    assert all(e.group_id == world.elv.id and e.is_active for e in rows)
    assert rows[1].description == "main lift" and rows[0].description is None and rows[2].description is None
    audits = (await db.execute(select(m.AuditLog).where(m.AuditLog.action == "equipment.import"))).scalars().all()
    assert len(audits) == 1 and audits[0].entity_id == out["batch_id"] and audits[0].actor_id == world.boss.id
    assert audits[0].metadata_json["created"] == 3 and audits[0].metadata_json["file_sha256"] == sha(data)
    assert "SE-01-1.1" not in str(audits[0].metadata_json)  # codes are not copied into the audit row


async def test_only_valid_rows_are_imported_and_every_rejection_has_row_number_and_reason(api, world, db):
    data = make_xlsx([
        ["OK-1", "ELV", None],                  # row 2 valid
        ["BAD CODE", "ELV", None],              # row 3 space inside
        ["OK-2", "DOOR", None],                 # row 4 wrong group
        ["", "ELV", None],                      # row 5 blank code
        ["OK-3", "", None],                     # row 6 blank group
        ["ok-1", "ELV", None],                  # row 7 duplicate of row 2 (case-insensitive)
        ["E\u200c-9", "ELV", None],             # row 8 zero-width non-joiner (invisible!)
        ["=1+1", "ELV", None],                  # row 9 formula
        ["OK-4", "ELV", "=HYPERLINK(1)"],       # row 10 formula in description
        ["ELV-1", "ELV", "overwrite me"],       # row 11 already exists
        ["OK-5", "ELV", "x" * 2001],            # row 12 description too long
        ["OK-6", "XXX", None],                  # row 13 unknown group code
        ["OK-7", "ELV", None],                  # row 14 valid
    ])
    p = (await preview(api, "boss", world.elv.id, data)).json()
    assert (p["rows_total"], p["valid_count"], p["error_count"]) == (13, 2, 11)
    assert {e["row"]: e["code"] for e in p["errors"]} == {
        3: "invalid_equipment_code", 4: "group_mismatch", 5: "missing_equipment_code", 6: "missing_group_code",
        7: "duplicate_in_file", 8: "invalid_equipment_code", 9: "formula_not_allowed", 10: "formula_not_allowed",
        11: "already_exists", 12: "invalid_description", 13: "group_mismatch",
    }
    by_row = {e["row"]: e for e in p["errors"]}
    assert by_row[7]["first_row"] == 2
    assert by_row[8]["bad_chars"] == ["U+200C"]            # the invisible character is named
    assert by_row[4]["found"] == "DOOR" and by_row[4]["expected"] == "ELV"
    assert by_row[11]["existing_group_code"] == "ELV" and by_row[11]["existing_is_active"] is True
    assert by_row[3]["equipment_code"] == "BAD CODE"

    c = await commit(api, "boss", world.elv.id, data)
    assert c.status_code == 201 and c.json()["created_count"] == 2 and c.json()["error_count"] == 11
    codes = set((await db.execute(select(m.Equipment.equipment_code))).scalars())
    assert {"OK-1", "OK-7"} <= codes and not ({"OK-2", "OK-3", "OK-4", "OK-5", "OK-6"} & codes)
    batch = (await db.execute(select(m.EquipmentImportBatch))).scalar_one()
    assert (batch.created_count, batch.rejected_count) == (2, 11)


async def test_unknown_group_code_never_creates_a_group_and_wrong_group_file_is_fully_rejected(api, world, db):
    groups_before = await count(db, m.Group)
    data = make_xlsx([[f"D-{i}", "DOOR", None] for i in range(5)])   # a DOOR file uploaded into ELV
    p = (await preview(api, "boss", world.elv.id, data)).json()
    assert p["valid_count"] == 0 and p["error_count"] == 5
    assert {e["code"] for e in p["errors"]} == {"group_mismatch"}
    r = await commit(api, "boss", world.elv.id, data)
    assert r.status_code == 422 and r.json()["detail"] == {"code": "nothing_to_import", "error_count": 5}
    assert await count(db, m.Group) == groups_before and await count(db, m.EquipmentImportBatch) == 0


async def test_existing_code_in_any_group_is_never_overwritten(api, world, db):
    world.d1.description = "original"
    await db.commit()
    data = make_xlsx([["door-1", "ELV", "NEW TEXT"], ["ELV-1", "ELV", "NEW TEXT"], ["FRESH-1", "ELV", None]])
    p = (await preview(api, "boss", world.elv.id, data)).json()
    assert {e["row"]: (e["code"], e.get("existing_group_code")) for e in p["errors"]} == {
        2: ("already_exists", "DOOR"), 3: ("already_exists", "ELV")}
    await commit(api, "boss", world.elv.id, data)
    await db.refresh(world.d1)
    await db.refresh(world.e1)
    assert world.d1.description == "original" and world.d1.group_id == world.door.id
    assert world.e1.description is None


async def test_inactive_or_unknown_group_is_refused(api, world, db):
    data = make_xlsx([["Z-1", "ELV", None]])
    assert (await preview(api, "boss", 9999, data)).status_code == 404
    world.elv.is_active = False
    await db.commit()
    r = await preview(api, "boss", world.elv.id, data)
    assert r.status_code == 409 and r.json()["detail"]["code"] == "group_inactive"
    assert (await commit(api, "boss", world.elv.id, data)).status_code == 409


# ======================= file-level validation =======================

@pytest.mark.parametrize("data,code", [
    (make_xlsx([["A-1", "ELV"]], sheets=2), "multiple_sheets"),
    (make_xlsx([["A-1"]], header=["equipment_code"]), "missing_columns"),
    (make_xlsx([["A-1", "ELV"]], header=["equipment_code", "group_code", "group_code"]), "duplicate_columns"),
    (make_xlsx([], header=HEADER), "no_data_rows"),
    (make_xlsx([], header=None), "empty_file"),
    (b"PK\x03\x04 this is not really a zip", "file_corrupt"),
    (b"just text", "file_corrupt"),
])
async def test_file_level_rejections(api, world, data, code):
    r = await preview(api, "boss", world.elv.id, data)
    detail = r.json()["detail"]
    got = detail["code"] if isinstance(detail, dict) else detail
    assert r.status_code == 422 and got == code, r.text


async def test_unsupported_extensions_are_refused(api, world):
    data = make_xlsx([["A-1", "ELV"]])
    for name in ("a.xls", "a.xlsm", "a.txt", "a", "a.xlsx.exe", "a.ods"):
        r = await preview(api, "boss", world.elv.id, data, name=name)
        assert r.status_code == 422 and r.json()["detail"] == "unsupported_file_type", name


async def test_extra_columns_are_ignored_but_reported_and_blank_header_cells_are_fine(api, world):
    data = make_xlsx([["A-1", "ELV", None, "junk", "more"]], header=HEADER + [None, "Notes"])
    p = (await preview(api, "boss", world.elv.id, data)).json()
    assert p["valid_count"] == 1 and p["ignored_columns"] == ["Notes"]


async def test_headers_are_case_and_nbsp_tolerant(api, world):
    data = make_xlsx([["A-1", "ELV", None]], header=[f" Equipment_Code{NBSP}", "GROUP_CODE", "Description"])
    assert (await preview(api, "boss", world.elv.id, data)).json()["valid_count"] == 1


async def test_date_cell_is_rejected_because_excel_turns_codes_like_1_10_into_dates(api, world):
    from datetime import date
    data = make_xlsx([[date(2026, 10, 1), "ELV", None], ["OK-1", "ELV", None]])
    p = (await preview(api, "boss", world.elv.id, data)).json()
    assert p["errors"] == [{**p["errors"][0], "row": 2, "code": "date_cell", "column": "equipment_code"}]
    assert p["valid_count"] == 1


async def test_numeric_code_cells_are_accepted_as_text(api, world):
    data = make_xlsx([[1001, "ELV", None], [1002.0, "ELV", None]])
    p = (await preview(api, "boss", world.elv.id, data)).json()
    assert [s["equipment_code"] for s in p["sample"]] == ["1001", "1002"]


async def test_row_limit_and_blank_rows(api, world, monkeypatch):
    monkeypatch.setattr(get_settings(), "import_max_rows", 5)
    ok = make_xlsx([[f"A-{i}", "ELV", None] for i in range(5)] + [[None, None, None]] * 3)
    assert (await preview(api, "boss", world.elv.id, ok)).json()["rows_total"] == 5   # blank rows do not count
    too_many = make_xlsx([[f"A-{i}", "ELV", None] for i in range(6)])
    r = await preview(api, "boss", world.elv.id, too_many)
    assert r.status_code == 422 and r.json()["detail"] == {"code": "too_many_rows", "max_rows": 5}


# ======================= csv =======================

async def test_csv_utf8_bom_semicolon_and_persian_description(api, world):
    text = "\ufeffequipment_code;group_code;description\r\nA-1;ELV;پله برقی\r\nA-2;ELV;\r\n"
    r = await api.post("/equipment/import/preview", "boss", data={"group_id": str(world.elv.id)},
                       files={"file": ("a.csv", text.encode("utf-8"), "text/csv")})
    assert r.status_code == 200 and r.json()["valid_count"] == 2
    assert r.json()["sample"][0]["description"] == "پله برقی"


async def test_csv_in_legacy_encoding_is_refused_with_a_clear_code(api, world):
    raw = "equipment_code,group_code,description\r\nA-1,ELV,پله\r\n".encode("cp1256")
    r = await api.post("/equipment/import/preview", "boss", data={"group_id": str(world.elv.id)},
                       files={"file": ("a.csv", raw, "text/csv")})
    assert r.status_code == 422 and r.json()["detail"] == "file_encoding"


async def test_csv_with_nul_bytes_is_refused(api, world):
    r = await api.post("/equipment/import/preview", "boss", data={"group_id": str(world.elv.id)},
                       files={"file": ("a.csv", b"equipment_code,group_code\r\nA\x00-1,ELV\r\n", "text/csv")})
    assert r.status_code == 422 and r.json()["detail"] == "file_corrupt"


# ======================= hostile files =======================

def _zip_with(extra_name, payload, base=None):
    src = zipfile.ZipFile(io.BytesIO(base or make_xlsx([["A-1", "ELV"]])))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            z.writestr(info.filename, src.read(info.filename))
        z.writestr(extra_name, payload)
    return out.getvalue()


async def test_zip_bomb_is_refused_by_compression_ratio_alone(api, world, monkeypatch):
    monkeypatch.setattr(get_settings(), "import_max_uncompressed_bytes", 10**10)   # only the ratio guard is left
    bomb = _zip_with("xl/media/zeros.dat", b"\x00" * (30 * 1024 * 1024))
    r = await preview(api, "boss", world.elv.id, bomb)
    assert r.status_code == 422 and r.json()["detail"] == "file_unsafe"


async def test_uncompressed_total_limit_is_enforced_on_an_ordinary_workbook(api, world, monkeypatch):
    monkeypatch.setattr(get_settings(), "import_max_uncompressed_bytes", 1000)
    r = await preview(api, "boss", world.elv.id, make_xlsx([["A-1", "ELV"]]))
    assert r.status_code == 422 and r.json()["detail"] == "file_unsafe"


async def test_zip_bomb_is_refused_before_openpyxl_runs(api, world):
    bomb = _zip_with("xl/media/zeros.bin".replace(".bin", ".dat"), b"\x00" * (30 * 1024 * 1024))
    assert len(bomb) < get_settings().import_max_bytes     # tiny upload, huge when expanded
    r = await preview(api, "boss", world.elv.id, bomb)
    assert r.status_code == 422 and r.json()["detail"] == "file_unsafe"


async def test_ordinary_excel_printer_settings_member_is_accepted(api, world):
    """Excel writes xl/printerSettings/printerSettings1.bin into normal workbooks — refusing every .bin was a bug."""
    data = _zip_with("xl/printerSettings/printerSettings1.bin", b"\x00" * 200)
    r = await preview(api, "boss", world.elv.id, data)
    assert r.status_code == 200 and r.json()["valid_count"] == 1


@pytest.mark.parametrize("member", ["xl/activeX/activeX1.xml", "xl/embeddings/oleObject1.bin"])
async def test_active_content_members_are_refused(api, world, member):
    r = await preview(api, "boss", world.elv.id, _zip_with(member, b"x"))
    assert r.status_code == 422 and r.json()["detail"] == "file_unsafe"


async def test_ratio_guard_also_covers_members_well_below_one_megabyte(api, world, monkeypatch):
    monkeypatch.setattr(get_settings(), "import_max_uncompressed_bytes", 10**10)   # only the ratio guard is left
    r = await preview(api, "boss", world.elv.id, _zip_with("xl/media/zeros.dat", b"\x00" * (300 * 1024)))
    assert r.status_code == 422 and r.json()["detail"] == "file_unsafe"


async def test_macro_enabled_content_is_refused(api, world):
    r = await preview(api, "boss", world.elv.id, _zip_with("xl/vbaProject.bin", b"MACRO"))
    assert r.status_code == 422 and r.json()["detail"] == "file_unsafe"


async def test_xml_entity_bomb_is_refused(api, world):
    from openpyxl.xml import DEFUSEDXML
    assert DEFUSEDXML, "defusedxml must be installed: it is what stops billion-laughs XML in uploads"
    laughs = ('<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol"><!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">'
              '<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">]>'
              '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&lol3;</t></is></c></row></sheetData></worksheet>')
    base = make_xlsx([["A-1", "ELV"]])
    src, out = zipfile.ZipFile(io.BytesIO(base)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            z.writestr(info.filename, laughs if info.filename == "xl/worksheets/sheet1.xml" else src.read(info.filename))
    r = await preview(api, "boss", world.elv.id, out.getvalue())
    assert r.status_code == 422 and r.json()["detail"] in ("file_corrupt", "file_unsafe")


async def test_oversized_upload_is_refused_without_reading_it(api, world):
    big = b"x" * (get_settings().import_max_bytes + 200_000)
    r = await preview(api, "boss", world.elv.id, big)
    assert r.status_code == 413 and r.json()["detail"]["code"] == "file_too_large"


async def test_declared_length_over_the_limit_is_refused_before_the_body_is_parsed(api, world):
    """Proves the guard sits in front of multipart parsing: the body here is junk that would otherwise be a 422."""
    big = str(get_settings().import_max_bytes + 1_000_000)
    r = await api.post("/equipment/import/preview", "boss", content=b"--x--",
                       headers={"content-type": "multipart/form-data; boundary=x", "content-length": big})
    assert r.status_code == 413 and r.json()["detail"]["code"] == "file_too_large"


async def test_chunked_upload_without_content_length_is_refused(api, world):
    async def body():
        yield b"--x\r\n"
    r = await api.post("/equipment/import/preview", "boss", content=body(),
                       headers={"content-type": "multipart/form-data; boundary=x"})
    assert r.status_code == 411


# ======================= authorization =======================

@pytest.mark.parametrize("who", ["ali", "exp", "aud", "omid"])
async def test_only_management_may_use_any_import_endpoint(api, world, db, who):
    data = make_xlsx([["Z-1", "ELV", None]])
    calls = [
        await preview(api, who, world.elv.id, data),
        await commit(api, who, world.elv.id, data),
        await api.get("/equipment/import/template", who),
        await api.get("/equipment/import/batches", who),
        await api.post("/equipment/import/batches/1/undo", who),
    ]
    assert [c.status_code for c in calls] == [403] * 5, [c.text for c in calls]
    assert await count(db, m.EquipmentImportBatch) == 0


async def test_unauthenticated_requests_are_refused(api, world):
    data = make_xlsx([["Z-1", "ELV", None]])
    assert (await preview(api, None, world.elv.id, data)).status_code == 401
    assert (await api.get("/equipment/import/batches")).status_code == 401
    assert (await api.post("/equipment/import/batches/1/undo")).status_code == 401


# ======================= commit integrity =======================

async def test_commit_requires_the_previewed_file(api, world, db):
    data = make_xlsx([["A-1", "ELV", None]])
    other = make_xlsx([["A-1", "ELV", None], ["A-2", "ELV", None]])
    r = await commit(api, "boss", world.elv.id, other, digest=sha(data))
    assert r.status_code == 409 and r.json()["detail"] == "file_changed"
    assert (await api.post("/equipment/import/commit", "boss", data={"group_id": str(world.elv.id)},
                           files={"file": ("a.xlsx", data, XLSX)})).status_code == 422   # digest is mandatory
    assert await count(db, m.EquipmentImportBatch) == 0


async def test_race_on_unique_code_rolls_back_everything(api, world, db, monkeypatch):
    """If another request creates one of the codes between our check and our insert, the DB unique constraint
    stops us and NOTHING (no batch, no equipment, no audit) is persisted."""
    from app.services import equipment_import as svc

    async def blind(_db, _codes):
        return {}
    monkeypatch.setattr(svc, "_existing", blind)
    data = make_xlsx([["FRESH-1", "ELV", None], ["ELV-1", "ELV", None]])   # ELV-1 really exists
    equipment_before = await count(db, m.Equipment)
    r = await commit(api, "boss", world.elv.id, data)
    assert r.status_code == 409 and r.json()["detail"] == "import_conflict"
    db.expire_all()
    assert await count(db, m.Equipment) == equipment_before
    assert await count(db, m.EquipmentImportBatch) == 0
    assert (await db.execute(select(m.AuditLog).where(m.AuditLog.action == "equipment.import"))).first() is None


# ======================= batches & undo =======================

async def _import(api, group, rows, who="boss"):
    data = make_xlsx(rows)
    r = await commit(api, who, group.id, data)
    assert r.status_code == 201, r.text
    return r.json()["batch_id"], data


async def test_undo_deletes_the_whole_batch_and_keeps_the_record(api, world, db):
    batch_id, data = await _import(api, world.elv, [["U-1", "ELV", None], ["U-2", "ELV", None]])
    listing = (await api.get("/equipment/import/batches", "boss")).json()
    row = listing["items"][0]
    assert row["id"] == batch_id and row["can_undo"] is True and row["remaining_count"] == 2
    assert row["group_code"] == "ELV" and row["created_by_name"] == "Boss" and row["undone_at"] is None

    boss_id, elv_id = world.boss.id, world.elv.id
    r = await api.post(f"/equipment/import/batches/{batch_id}/undo", "boss")
    assert r.status_code == 200 and r.json() == {"batch_id": batch_id, "deleted_count": 2}
    db.expire_all()
    assert (await db.execute(select(m.Equipment).where(m.Equipment.equipment_code.in_(["U-1", "U-2"])))).first() is None
    batch = (await db.execute(select(m.EquipmentImportBatch))).scalar_one()
    assert batch.undone_at is not None and batch.undone_by == boss_id
    audit = (await db.execute(select(m.AuditLog).where(m.AuditLog.action == "equipment.import_undo"))).scalar_one()
    assert audit.metadata_json["codes"] == ["U-1", "U-2"] and audit.entity_id == batch_id
    after = (await api.get("/equipment/import/batches", "boss")).json()["items"][0]
    assert after["can_undo"] is False and after["undone_at"] is not None

    assert (await api.post(f"/equipment/import/batches/{batch_id}/undo", "boss")).json()["detail"] == "batch_already_undone"
    again = await commit(api, "boss", elv_id, data)       # the corrected re-import works: codes are free again
    assert again.status_code == 201 and again.json()["created_count"] == 2


async def test_undo_is_refused_once_any_equipment_has_a_report(api, world, db):
    batch_id, _ = await _import(api, world.elv, [["R-1", "ELV", None], ["R-2", "ELV", None]])
    r1 = (await db.execute(select(m.Equipment).where(m.Equipment.equipment_code == "R-1"))).scalar_one()
    await f.report(db, "05-ELV-001", r1, world.elv, world.ali)
    await db.commit()
    r = await api.post(f"/equipment/import/batches/{batch_id}/undo", "boss")
    assert r.status_code == 409
    assert r.json()["detail"] == {"code": "batch_has_reports", "report_count": 1, "equipment_codes": ["R-1"]}
    db.expire_all()
    assert (await db.execute(select(func.count()).select_from(m.Equipment)
                             .where(m.Equipment.import_batch_id == batch_id))).scalar_one() == 2    # nothing deleted
    assert (await api.get("/equipment/import/batches", "boss")).json()["items"][0]["can_undo"] is False


async def test_soft_deleted_reports_also_block_undo(api, world, db):
    batch_id, _ = await _import(api, world.elv, [["S-1", "ELV", None]])
    e = (await db.execute(select(m.Equipment).where(m.Equipment.equipment_code == "S-1"))).scalar_one()
    rep = await f.report(db, "05-ELV-002", e, world.elv, world.ali)
    rep.is_active = False
    await db.commit()
    assert (await api.post(f"/equipment/import/batches/{batch_id}/undo", "boss")).json()["detail"]["code"] == "batch_has_reports"


async def test_undo_is_refused_when_equipment_was_moved_to_another_group(api, world, db):
    batch_id, _ = await _import(api, world.elv, [["M-1", "ELV", None], ["M-2", "ELV", None]])
    e = (await db.execute(select(m.Equipment).where(m.Equipment.equipment_code == "M-1"))).scalar_one()
    assert (await api.post(f"/equipment/{e.id}/move", "boss", json={"group_id": world.door.id})).status_code == 200
    r = await api.post(f"/equipment/import/batches/{batch_id}/undo", "boss")
    assert r.status_code == 409 and r.json()["detail"] == {"code": "batch_equipment_moved", "moved_count": 1}


async def test_undo_race_path_has_the_same_response_contract(api, world, db, monkeypatch):
    """The guard misses (a report appears between check and delete); the FK stops the delete. The answer must carry
    the same fields as the normal refusal."""
    from app.services import equipment_import as svc
    batch_id, _ = await _import(api, world.elv, [["Q-1", "ELV", None]])
    e = (await db.execute(select(m.Equipment).where(m.Equipment.equipment_code == "Q-1"))).scalar_one()
    await f.report(db, "05-ELV-003", e, world.elv, world.ali)
    await db.commit()
    real = svc._batch_stats
    calls = {"n": 0}

    async def blind_first_time(_db, ids):
        calls["n"] += 1
        stats = await real(_db, ids)
        if calls["n"] == 1:
            for s in stats.values():
                s["reports"] = 0
        return stats
    monkeypatch.setattr(svc, "_batch_stats", blind_first_time)
    r = await api.post(f"/equipment/import/batches/{batch_id}/undo", "boss")
    assert r.status_code == 409
    assert r.json()["detail"] == {"code": "batch_has_reports", "report_count": 1, "equipment_codes": ["Q-1"]}


async def test_moving_equipment_out_and_back_does_not_restore_undo(api, world, db):
    """CodeRabbit finding: comparing only the CURRENT group let A -> B -> A look 'never moved'."""
    batch_id, _ = await _import(api, world.elv, [["W-1", "ELV", None], ["W-2", "ELV", None]])
    e = (await db.execute(select(m.Equipment).where(m.Equipment.equipment_code == "W-1"))).scalar_one()
    eid, elv_id, door_id = e.id, world.elv.id, world.door.id
    assert (await api.post(f"/equipment/{eid}/move", "boss", json={"group_id": door_id})).status_code == 200
    assert (await api.post(f"/equipment/{eid}/move", "boss", json={"group_id": elv_id})).status_code == 200
    listing = (await api.get("/equipment/import/batches", "boss")).json()["items"][0]
    assert listing["can_undo"] is False
    r = await api.post(f"/equipment/import/batches/{batch_id}/undo", "boss")
    assert r.status_code == 409 and r.json()["detail"] == {"code": "batch_equipment_moved", "moved_count": 1}
    db.expire_all()
    assert (await db.execute(select(func.count()).select_from(m.Equipment)
                             .where(m.Equipment.import_batch_id == batch_id))).scalar_one() == 2


async def test_undo_unknown_batch_is_404_and_manual_equipment_is_never_touched(api, world, db):
    assert (await api.post("/equipment/import/batches/9999/undo", "boss")).status_code == 404
    batch_id, _ = await _import(api, world.elv, [["N-1", "ELV", None]])
    await api.post(f"/equipment/import/batches/{batch_id}/undo", "boss")
    db.expire_all()
    codes = set((await db.execute(select(m.Equipment.equipment_code))).scalars())
    assert {"ELV-1", "DOOR-1"} <= codes        # created outside the import: untouched


async def test_batches_list_pagination_and_newest_first(api, world):
    first, _ = await _import(api, world.elv, [["P-1", "ELV", None]])
    second, _ = await _import(api, world.elv, [["P-2", "ELV", None]])
    page = (await api.get("/equipment/import/batches?limit=1", "boss")).json()
    assert page["total"] == 2 and [i["id"] for i in page["items"]] == [second]
    assert (await api.get("/equipment/import/batches?limit=1&offset=1", "boss")).json()["items"][0]["id"] == first
    assert (await api.get("/equipment/import/batches?limit=0", "boss")).status_code == 422


# ======================= template =======================

async def test_xlsx_template_is_header_only_text_formatted_and_round_trips(api, world):
    r = await api.get("/equipment/import/template", "boss")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert "equipment-import-template.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == HEADER and ws.max_row >= 2
    assert all(ws.cell(row=2, column=c).number_format == "@" for c in (1, 2, 3))
    assert all(c.value is None for row in ws.iter_rows(min_row=2, max_row=5) for c in row)   # no sample data
    empty = await preview(api, "boss", world.elv.id, r.content)
    assert empty.status_code == 422 and empty.json()["detail"] == "no_data_rows"
    ws.append(["T-1", "ELV", "from template"])
    buf = io.BytesIO()
    wb.save(buf)
    ok = await preview(api, "boss", world.elv.id, buf.getvalue())
    assert ok.status_code == 200 and ok.json()["valid_count"] == 1


async def test_csv_template(api, world):
    r = await api.get("/equipment/import/template?format=csv", "boss")
    assert r.status_code == 200 and r.content.decode("utf-8-sig").strip() == "equipment_code,group_code,description"
    assert (await api.get("/equipment/import/template?format=exe", "boss")).status_code == 422
